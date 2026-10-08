"""
generate_samples.py — Create the two binary sample documents:
  - sample_docs/insurance_policy.docx   (python-docx)
  - sample_docs/it_faq.pdf              (reportlab)

Run once from the project root:
    python generate_samples.py
"""

from pathlib import Path

# ── Insurance Policy DOCX ─────────────────────────────────────────────────────
def create_insurance_docx():
    import docx
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = docx.Document()

    # Title
    title = doc.add_heading("TrustLife Group Health Insurance Policy", level=1)
    title.runs[0].font.color.rgb = RGBColor(0x1f, 0x69, 0xeb)

    doc.add_paragraph("Policy Number: TL-GHI-2024-7890")
    doc.add_paragraph("Effective Date: 1 January 2024")
    doc.add_paragraph("Policyholder: ACME Corporation")
    doc.add_paragraph("")

    # Section 1
    doc.add_heading("Section 1: Overview", level=2)
    doc.add_paragraph(
        "This Group Health Insurance Policy ('Policy') is issued by TrustLife Insurance "
        "Limited to ACME Corporation ('Employer') for the benefit of eligible employees "
        "and their registered dependants. The Policy is governed by the Insurance Act 2015 "
        "and the terms and conditions set out herein."
    )

    # Section 2
    doc.add_heading("Section 2: Coverage Limits", level=2)
    doc.add_paragraph(
        "2.1 Annual Coverage Limit\n"
        "Each insured employee is entitled to health insurance coverage up to a maximum of "
        "INR 500,000 (Indian Rupees Five Lakhs) per policy year. This limit applies to the "
        "aggregate of all eligible medical expenses incurred by the employee and registered "
        "dependants within the same policy year."
    )
    doc.add_paragraph(
        "2.2 Individual Sub-limits\n"
        "The following sub-limits apply within the overall INR 500,000 limit:\n"
        "- Hospitalisation (room & board): INR 5,000 per day, up to 30 days per year.\n"
        "- ICU charges: INR 10,000 per day, up to 15 days per year.\n"
        "- Day-care procedures: INR 50,000 per procedure.\n"
        "- Pre- and post-hospitalisation: 30 days before and 60 days after discharge.\n"
        "- Ambulance charges: INR 2,000 per event."
    )
    doc.add_paragraph(
        "2.3 Outpatient Benefits\n"
        "Outpatient consultations and diagnostics are covered up to INR 15,000 per year "
        "per insured individual, subject to a copayment of 20% per visit."
    )

    # Section 3
    doc.add_heading("Section 3: Exclusions", level=2)
    doc.add_paragraph(
        "The following are not covered under this Policy:\n"
        "- Cosmetic or aesthetic treatments.\n"
        "- Self-inflicted injuries.\n"
        "- Experimental or unproven treatments.\n"
        "- Dental treatment (unless arising from an accident).\n"
        "- Vision correction (spectacles, contact lenses, LASIK).\n"
        "- Pre-existing conditions not disclosed at enrolment (first 12 months only).\n"
        "- Substance abuse or addiction treatment."
    )

    # Section 4
    doc.add_heading("Section 4: Claims Procedure", level=2)
    doc.add_paragraph(
        "4.1 Cashless Claims\n"
        "Employees may avail cashless treatment at any of TrustLife's 3,000+ empanelled "
        "network hospitals. Present your TrustLife Health Card at the hospital's insurance "
        "desk. Pre-authorisation is required for planned admissions at least 48 hours in "
        "advance; for emergency admissions, notify TrustLife within 24 hours."
    )
    doc.add_paragraph(
        "4.2 Reimbursement Claims\n"
        "For treatment at non-network hospitals, the employee pays upfront and submits "
        "original bills, discharge summary, and claim form to HR within 30 days of discharge. "
        "Reimbursement is processed within 15 working days of receipt of complete documents."
    )

    # Section 5
    doc.add_heading("Section 5: Premium and Renewal", level=2)
    doc.add_paragraph(
        "The annual premium is fully borne by the Employer. The Policy renews automatically "
        "on 1 January each year subject to mutual agreement. Premium revision, if any, will "
        "be notified at least 60 days before renewal."
    )

    # Section 6
    doc.add_heading("Section 6: Contact Information", level=2)
    doc.add_paragraph(
        "For claims assistance: 1800-TL-CLAIM (toll-free, 24x7)\n"
        "Email: claims@trustlife.example.com\n"
        "Website: www.trustlife.example.com/employee-portal"
    )

    out_path = Path("sample_docs/insurance_policy.docx")
    doc.save(str(out_path))
    print(f"Created {out_path}")


# ── IT FAQ PDF ────────────────────────────────────────────────────────────────
def create_it_faq_pdf():
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
    from reportlab.lib.enums import TA_LEFT

    out_path = Path("sample_docs/it_faq.pdf")
    doc = SimpleDocTemplate(
        str(out_path),
        pagesize=A4,
        leftMargin=2.5*cm, rightMargin=2.5*cm,
        topMargin=2.5*cm, bottomMargin=2.5*cm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "Title2", parent=styles["Title"],
        textColor=colors.HexColor("#1f69eb"), fontSize=20, spaceAfter=12,
    )
    h2_style = ParagraphStyle(
        "H2", parent=styles["Heading2"],
        textColor=colors.HexColor("#0d47a1"), fontSize=13, spaceBefore=16, spaceAfter=4,
    )
    body_style = ParagraphStyle(
        "Body2", parent=styles["Normal"],
        fontSize=10, leading=16, spaceAfter=6,
    )
    bullet_style = ParagraphStyle(
        "Bullet2", parent=styles["Normal"],
        fontSize=10, leading=16, leftIndent=20, spaceAfter=4,
        bulletText="•",
    )

    story = []

    story.append(Paragraph("ACME Corporation — IT Help Desk FAQ", title_style))
    story.append(Paragraph("Version 2.1 | Updated: March 2024", body_style))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#cccccc")))
    story.append(Spacer(1, 0.3*cm))

    story.append(Paragraph("Section 1: Account & Access", h2_style))

    story.append(Paragraph("<b>Q1: How do I reset my password?</b>", body_style))
    story.append(Paragraph(
        "If you have forgotten your password or need to reset it, follow these steps:", body_style))
    steps = [
        "Step 1: Navigate to the ACME Self-Service Portal at https://itsupport.acmecorp.example.com",
        "Step 2: Click the <b>'Forgot Password'</b> link on the login page.",
        "Step 3: Enter your registered corporate email address and click <b>'Send Reset Link'</b>.",
        "Step 4: Check your email inbox (and spam/junk folder) for a reset email from noreply@acmecorp.example.com. The link expires in 30 minutes.",
        "Step 5: Click the link in the email and enter your new password. Passwords must be at least 12 characters and include uppercase, lowercase, a number, and a special character.",
        "Step 6: Log in to your account using the new password. You will be prompted to set up MFA if not already configured.",
    ]
    for step in steps:
        story.append(Paragraph(step, bullet_style))
    story.append(Paragraph(
        "If you do not receive the email within 5 minutes, contact the IT Help Desk at ext. 1001 or helpdesk@acmecorp.example.com.",
        body_style
    ))
    story.append(Spacer(1, 0.2*cm))

    story.append(Paragraph("<b>Q2: How do I unlock my Active Directory account?</b>", body_style))
    story.append(Paragraph(
        "After 5 consecutive failed login attempts, your account will be locked for 15 minutes automatically. "
        "To unlock immediately, call the IT Help Desk at extension 1001 with your employee ID for verification. "
        "Alternatively, your account will unlock automatically after 15 minutes of inactivity.",
        body_style
    ))
    story.append(Spacer(1, 0.2*cm))

    story.append(Paragraph("<b>Q3: How do I set up Multi-Factor Authentication (MFA)?</b>", body_style))
    story.append(Paragraph(
        "MFA is mandatory for all employees. To set it up:",
        body_style
    ))
    mfa_steps = [
        "Download the Microsoft Authenticator app on your smartphone.",
        "Log in to the ACME Office 365 portal at https://portal.acmecorp.example.com.",
        "Follow the on-screen prompts to scan the QR code with the Authenticator app.",
        "Approve the test notification to complete setup.",
    ]
    for s in mfa_steps:
        story.append(Paragraph(s, bullet_style))
    story.append(Spacer(1, 0.3*cm))

    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#eeeeee")))
    story.append(Paragraph("Section 2: Hardware & Equipment", h2_style))

    story.append(Paragraph("<b>Q4: How do I request a new laptop or peripheral?</b>", body_style))
    story.append(Paragraph(
        "Hardware requests must be submitted through the IT Portal (https://itsupport.acmecorp.example.com) "
        "under 'New Equipment Request'. Requests require manager approval. Standard delivery time is 5-7 working days. "
        "For urgent requests, contact IT Help Desk directly.",
        body_style
    ))
    story.append(Spacer(1, 0.2*cm))

    story.append(Paragraph("<b>Q5: What do I do if my laptop is slow or unresponsive?</b>", body_style))
    steps_laptop = [
        "Restart the laptop — this resolves most performance issues.",
        "Check for pending Windows/macOS updates and install them.",
        "Close unused applications and browser tabs.",
        "Run a virus scan using the pre-installed CrowdStrike Falcon agent.",
        "If the issue persists, submit an IT ticket with a description of the problem.",
    ]
    for s in steps_laptop:
        story.append(Paragraph(s, bullet_style))
    story.append(Spacer(1, 0.3*cm))

    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#eeeeee")))
    story.append(Paragraph("Section 3: Software & Applications", h2_style))

    story.append(Paragraph("<b>Q6: How do I install approved software?</b>", body_style))
    story.append(Paragraph(
        "Approved software is available in the Company App Store (Software Center on Windows, "
        "Self-Service on macOS). Browse, select, and click Install. No admin password is required. "
        "Non-approved software must not be installed without a formal exception request approved by "
        "your manager and the IT Security team.",
        body_style
    ))
    story.append(Spacer(1, 0.2*cm))

    story.append(Paragraph("<b>Q7: How do I access company resources from home (VPN)?</b>", body_style))
    story.append(Paragraph(
        "Use the Cisco AnyConnect VPN client pre-installed on your laptop. Connect to "
        "vpn.acmecorp.example.com using your corporate email and password, then approve the MFA prompt. "
        "Disconnect VPN when accessing only internet resources to preserve bandwidth.",
        body_style
    ))
    story.append(Spacer(1, 0.3*cm))

    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#eeeeee")))
    story.append(Paragraph("Section 4: IT Policies", h2_style))

    story.append(Paragraph("<b>Q8: What is the Acceptable Use Policy (AUP)?</b>", body_style))
    story.append(Paragraph(
        "Company IT resources (laptops, email, internet) are provided for business use. "
        "Limited personal use is tolerated but must not interfere with work, consume excessive bandwidth, "
        "or involve illegal content. All activity on company devices is subject to monitoring. "
        "The full AUP is available on the intranet at https://intranet.acmecorp.example.com/policies.",
        body_style
    ))
    story.append(Spacer(1, 0.2*cm))

    story.append(Paragraph("<b>Q9: How should I report a security incident or phishing email?</b>", body_style))
    story.append(Paragraph(
        "Forward any suspicious email to security@acmecorp.example.com immediately. Do not click links "
        "or download attachments from unknown senders. For active security incidents (ransomware, data breach), "
        "call the Security Hotline at ext. 9911 immediately and disconnect your device from the network.",
        body_style
    ))
    story.append(Spacer(1, 0.3*cm))

    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#cccccc")))
    story.append(Spacer(1, 0.2*cm))
    story.append(Paragraph(
        "IT Help Desk: ext. 1001 | helpdesk@acmecorp.example.com | Mon–Fri 08:00–20:00",
        body_style
    ))

    doc.build(story)
    print(f"Created {out_path}")


if __name__ == "__main__":
    create_insurance_docx()
    create_it_faq_pdf()
    print("All sample documents created successfully.")
